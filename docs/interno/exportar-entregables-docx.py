"""Genera los .docx de docs/entregables a partir de los Markdown."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1] / "entregables"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x55, 0x55, 0x55)

FILES = [
    ("Alcance-funcional.md", "Alcance funcional.docx"),
    ("Supuestos-y-decisiones.md", "Supuestos y decisiones de negocio.docx"),
    ("Requerimientos-funcionales.md", "Requerimientos funcionales.docx"),
    ("Historias-de-usuario.md", "Historias de usuario.docx"),
    ("Incrementos-de-los-sprints.md", "Incrementos de los sprints.docx"),
    ("Modelo-de-datos.md", "Modelo de datos.docx"),
    ("Diseno-de-APIs.md", "Diseño de APIs.docx"),
    ("Testing-y-automatizacion.md", "Testing y automatización.docx"),
]

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\n]+\*)")


def set_run_font(run, name, size, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = name
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        r_fonts.set(qn(attr), name)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_cell_margins(cell, margin=40):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(margin))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def add_inline(paragraph, text, size=11, bold=False, color=None, font="Calibri"):
    for part in INLINE.split(text):
        if part == "":
            continue
        run = paragraph.add_run()
        part_bold = bold
        part_italic = False
        part_font = font
        part_size = size
        part_color = color
        value = part
        if part.startswith("**") and part.endswith("**"):
            value = part[2:-2]
            part_bold = True
        elif part.startswith("`") and part.endswith("`"):
            value = part[1:-1]
            part_font = "Consolas"
            part_size = max(8, size - 1)
        elif part.startswith("*") and part.endswith("*"):
            value = part[1:-1]
            part_italic = True
        set_run_font(run, part_font, part_size, part_bold, part_italic, part_color)
        run.text = value


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, "Calibri", 9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.08

    sizes = {1: 18, 2: 14, 3: 12}
    for level, size in sizes.items():
        style = document.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.paragraph_format.space_before = Pt(14 if level == 1 else 12)
        style.paragraph_format.space_after = Pt(6)


def configure_section(section, title):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("Mordi  ·  Desarrollo de aplicaciones  ·  2.º cuatrimestre 2026")
    set_run_font(run, "Calibri", 9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left = footer.add_run(f"{title}   ·   ")
    set_run_font(left, "Calibri", 9, color=MUTED)
    add_page_number(footer)


def split_row(line):
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    return [cell.strip() for cell in text.split("|")]


def is_separator(line):
    cells = split_row(line)
    if not cells:
        return False
    return all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def content_width_cm():
    return 21.0 - 1.6 - 1.6


def add_table(document, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    rows = [row + [""] * (cols - len(row)) for row in rows]
    table = document.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = False

    weights = []
    for col in range(cols):
        longest = max(len(row[col]) for row in rows)
        weights.append(max(longest, 6))
    total = sum(weights)
    width = content_width_cm()
    widths = [max(1.5, width * weight / total) for weight in weights]
    scale = width / sum(widths)
    widths = [item * scale for item in widths]

    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.rows[r_index].cells[c_index]
            cell.text = ""
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            header = r_index == 0
            add_inline(
                paragraph,
                value,
                size=9,
                bold=header,
                color=RGBColor(0xFF, 0xFF, 0xFF) if header else None,
            )
            if header:
                shade_cell(cell, "1F4E79")
            elif r_index % 2 == 0:
                shade_cell(cell, "F4F7FB")
            cell.width = Cm(widths[c_index])
        table.rows[r_index].cells[0].width = Cm(widths[0])

    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    tr_pr.append(header_flag)

    grid = table._tbl.tblGrid
    for index, child in enumerate(list(grid)):
        child.set(qn("w:w"), str(int(widths[index] * 567)))

    document.add_paragraph()


def add_code(document, code, language):
    if language == "mermaid":
        add_mermaid(document, code)
        return
    label = document.add_paragraph()
    label.paragraph_format.space_after = Pt(2)
    caption = label.add_run(language or "texto")
    set_run_font(caption, "Calibri", 9, italic=True, color=MUTED)
    for line in code.splitlines() or [""]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        shade_paragraph(paragraph, "F4F4F4")
        run = paragraph.add_run(line if line else " ")
        set_run_font(run, "Consolas", 8)
    document.add_paragraph()


def add_mermaid(document, code):
    intro = document.add_paragraph()
    add_inline(
        intro,
        "Entidades del diagrama (la cardinalidad está en la tabla que sigue).",
        size=11,
        color=MUTED,
    )
    current = None
    entities = []
    for line in code.splitlines():
        match = re.match(r"\s*([A-Za-z0-9_]+)\s*\{", line)
        if match:
            current = [match.group(1), []]
            continue
        if current and line.strip() == "}":
            entities.append(current)
            current = None
            continue
        if current and line.strip():
            current[1].append(line.strip())
    for name, fields in entities:
        heading = document.add_paragraph(style="Heading 3")
        add_inline(heading, name, size=12, bold=True, color=NAVY)
        parsed = [["Campo", "Tipo", "Notas"]]
        for field in fields:
            parts = field.split()
            if len(parts) >= 2:
                parsed.append([parts[1], parts[0], " ".join(parts[2:])])
            else:
                parsed.append([field, "", ""])
        add_table(document, parsed)


def add_body_paragraph(document, text, bullet=False, number=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    if bullet or number is not None:
        paragraph.paragraph_format.left_indent = Cm(0.6)
        prefix = "•  " if bullet else f"{number}.  "
        run = paragraph.add_run(prefix)
        set_run_font(run, "Calibri", 11)
    add_inline(paragraph, text, size=11)
    return paragraph


def convert(markdown, title):
    document = Document()
    configure_styles(document)
    configure_section(document.sections[0], title)
    document.core_properties.title = title
    document.core_properties.subject = "Carpeta de entregables — Mordi"
    document.core_properties.category = "Desarrollo de aplicaciones"

    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            language = stripped[3:].strip().lower()
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            add_code(document, "\n".join(block), language)
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_separator(lines[index + 1]):
            rows = [split_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index]))
                index += 1
            add_table(document, rows)
            continue

        if stripped in ("---", "***"):
            index += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_paragraph(style=f"Heading {level}")
            add_inline(paragraph, heading.group(2), size={1: 18, 2: 14, 3: 12}[level], bold=True, color=NAVY)
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            add_body_paragraph(document, bullet.group(1), bullet=True)
            index += 1
            continue

        number = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if number:
            add_body_paragraph(document, number.group(2), number=number.group(1))
            index += 1
            continue

        if stripped == "":
            index += 1
            continue

        add_body_paragraph(document, stripped)
        index += 1

    return document


def main():
    for source_name, target_name in FILES:
        source = ROOT / source_name
        markdown = source.read_text(encoding="utf-8")
        title = target_name.replace(".docx", "")
        document = convert(markdown, title)
        target = ROOT / target_name
        document.save(target)
        print(f"{target.name}: {target.stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
